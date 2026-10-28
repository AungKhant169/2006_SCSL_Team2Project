"""Build the Word report from Markdown. Requires python-docx (pip install python-docx)."""

from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def plain(text):
    return re.sub(r'`([^`]+)`', r'\1', text).replace('**', '')


def main():
    folder = Path(__file__).resolve().parent
    source = folder / 'Lab4_Coding_Agent_Report.md'
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(1.8)

    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name in ['Title', 'Heading 1', 'Heading 2']:
        document.styles[name].font.color.rgb = RGBColor.from_string('17365D')

    document.core_properties.title = 'Lab 4: Coding Agent Exercises — PathSG'
    document.core_properties.subject = 'Sections 3.3.2, 3.3.3 and 3.3.4'
    document.core_properties.author = '2006 SCSL Team 2'
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('PathSG | Lab 4 | ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    lines = source.read_text(encoding='utf-8').splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                cells = [plain(cell.strip()) for cell in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+', cell) for cell in cells):
                    rows.append(cells)
                index += 1
            table = document.add_table(rows=1, cols=len(rows[0]))
            table.style = 'Light Shading Accent 1'
            for column, value in enumerate(rows[0]):
                table.rows[0].cells[column].text = value
            header = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(header)
            for values in rows[1:]:
                cells = table.add_row().cells
                for column, value in enumerate(values):
                    cells[column].text = value
            for row in table.rows:
                no_split = OxmlElement('w:cantSplit')
                row._tr.get_or_add_trPr().append(no_split)
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        for run in paragraph.runs:
                            run.font.size = Pt(8)
            document.add_paragraph()
            continue
        if line.startswith('# '):
            document.add_heading(plain(line[2:]), 0)
        elif line.startswith('## '):
            document.add_heading(plain(line[3:]), 1)
        elif line.startswith('### '):
            document.add_heading(plain(line[4:]), 2)
        elif line.startswith('> '):
            document.add_paragraph(plain(line[2:]), style='Intense Quote')
        elif re.match(r'^\d+\. ', line):
            document.add_paragraph(plain(re.sub(r'^\d+\. ', '', line)), style='List Number')
        else:
            document.add_paragraph(plain(line))
        index += 1

    output = folder / 'Lab4_Coding_Agent_Report.docx'
    document.save(output)
    # Keep font metadata English-only as well as the visible report text.
    with ZipFile(output) as archive:
        entries = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
    with ZipFile(output, 'w') as archive:
        for entry, data in entries:
            if entry.filename in ('word/fontTable.xml', 'word/theme/theme1.xml'):
                data = re.sub(
                    r'="[^"<>]*[\u3400-\u9fff][^"<>]*"',
                    '="Calibri"',
                    data.decode('utf-8'),
                ).encode('utf-8')
            archive.writestr(entry, data)
    reopened = Document(output)
    assert len(reopened.tables) == 2
    assert any(p.text == '3.3.4 Reflections' for p in reopened.paragraphs)
    assert any('620 frontend tests' in p.text for p in reopened.paragraphs)
    print(f'Created and reopened {output.name}: {len(reopened.paragraphs)} paragraphs, 2 tables.')


if __name__ == '__main__':
    main()
